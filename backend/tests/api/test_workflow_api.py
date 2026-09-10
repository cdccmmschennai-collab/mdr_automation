"""upload -> extract -> automate -> summary over HTTP, on the small fixture.

Every test here goes through the real routes, the real services, the real
engine and the real PostgreSQL schema; only the workbook is small. What the
database ends up holding is compared with what `run_automation` produces for
the same file, because the whole point of Delivery Phase 3 is that the stored
result *is* the engine's result and not a re-derivation of it.

The real ~22k-row workbook goes through the same path in
`test_workflow_real_workbook.py`.
"""

from __future__ import annotations

import dataclasses
import hashlib
import uuid

import pytest
from openpyxl import load_workbook

from app.domain.enums.lifecycle import SubmissionStatus
from app.services import workflow_service
from app.services.automation_service import run_automation
from app.services.rule_set_service import file_digest
from tests.support.automation import DOCUMENT_ROWS
from tests.support.workbook import RULES_WORKBOOK, requires_rules_workbook

from .conftest import stored_rows, stored_submission, upload

SHA_HEX = 64


# ============================================================ UPLOAD

class TestUpload:

    def test_a_valid_workbook_creates_a_submission(self, client, plant,
                                                   small_workbook):
        response = upload(client, plant, small_workbook)
        assert response.status_code == 201, response.text
        body = response.json()
        assert uuid.UUID(body["mdr_id"])
        assert body["status"] == SubmissionStatus.UPLOADED
        assert body["plant_id"] == str(plant.id)
        assert body["source_filename"] == small_workbook.name
        assert body["submission_no"] >= 1

    def test_the_digest_and_size_are_of_the_bytes_received(self, client, plant,
                                                          small_workbook):
        data = small_workbook.read_bytes()
        body = upload(client, plant, small_workbook).json()
        assert body["source_sha256"] == hashlib.sha256(data).hexdigest()
        assert len(body["source_sha256"]) == SHA_HEX
        assert body["source_byte_size"] == len(data)

    def test_the_submission_is_persisted(self, client, plant, small_workbook,
                                         workflow_settings):
        body = upload(client, plant, small_workbook).json()
        stored = stored_submission(workflow_settings, body["mdr_id"])
        assert stored is not None
        assert stored.plant_id == plant.id
        assert stored.status == SubmissionStatus.UPLOADED
        assert stored.source_sha256 == body["source_sha256"]
        assert stored.failure_reason == ""

    def test_the_workbook_is_stored_under_the_submission(self, client, plant,
                                                         small_workbook,
                                                         workflow_settings):
        """`<uploads_dir>/<mdr_id>/<filename>`, byte for byte, and the key -
        not the absolute path - is what the row records."""
        body = upload(client, plant, small_workbook).json()
        stored = stored_submission(workflow_settings, body["mdr_id"])
        assert stored.stored_path == f"{body['mdr_id']}/{small_workbook.name}"

        path = workflow_settings.uploads_dir / stored.stored_path
        assert path.is_file()
        assert path.read_bytes() == small_workbook.read_bytes()
        assert path.parent.name == body["mdr_id"]
        assert workflow_settings.uploads_dir == workflow_settings.uploads_dir_override

    def test_submission_numbers_count_up_within_the_plant(self, client, plant,
                                                          small_workbook):
        first = upload(client, plant, small_workbook).json()
        second = upload(client, plant, small_workbook).json()
        assert second["submission_no"] == first["submission_no"] + 1
        assert second["mdr_id"] != first["mdr_id"]

    def test_the_same_bytes_may_be_uploaded_twice_without_overwriting(
            self, client, plant, small_workbook, workflow_settings):
        """Two submissions, two stored files. Nothing replaced."""
        a = upload(client, plant, small_workbook).json()
        b = upload(client, plant, small_workbook).json()
        a_path = workflow_settings.uploads_dir / f"{a['mdr_id']}/{small_workbook.name}"
        b_path = workflow_settings.uploads_dir / f"{b['mdr_id']}/{small_workbook.name}"
        assert a_path.is_file() and b_path.is_file() and a_path != b_path

    def test_an_unknown_plant_is_refused_and_nothing_is_stored(
            self, client, small_workbook, workflow_settings):
        fake_plant = type("P", (), {"id": uuid.uuid4()})()
        before = set(workflow_settings.uploads_dir.iterdir())
        response = upload(client, fake_plant, small_workbook)
        assert response.status_code == 404
        assert "plant" in response.json()["detail"].lower()
        assert set(workflow_settings.uploads_dir.iterdir()) == before

    def test_an_empty_file_is_refused(self, client, plant, tmp_path):
        empty = tmp_path / "empty.xlsx"
        empty.write_bytes(b"")
        response = upload(client, plant, empty)
        assert response.status_code == 400
        assert "empty" in response.json()["detail"].lower()

    def test_a_non_workbook_is_refused(self, client, plant, tmp_path):
        junk = tmp_path / "junk.xlsx"
        junk.write_bytes(b"PK\x03\x04 this is not a workbook")
        response = upload(client, plant, junk)
        assert response.status_code == 400
        assert "workbook" in response.json()["detail"].lower()

    def test_an_unsupported_extension_is_refused(self, client, plant,
                                                 small_workbook):
        response = upload(client, plant, small_workbook, filename="log.csv")
        assert response.status_code == 400
        assert ".xlsx" in response.json()["detail"]

    def test_a_workbook_without_the_mdr_sheet_is_refused(self, client, plant,
                                                         tmp_path):
        """A real xlsx, but not an MDR: refused before a submission exists."""
        from openpyxl import Workbook
        wb = Workbook()
        wb.active.title = "Sheet1"
        wb.active["A1"] = "hello"
        path = tmp_path / "not-mdr.xlsx"
        wb.save(path)
        response = upload(client, plant, path)
        assert response.status_code == 400
        assert "QatarEnergy" in response.json()["detail"]

    def test_an_oversized_upload_is_refused_over_http(self, client, plant,
                                                      small_workbook,
                                                      workflow_settings):
        from app.api.deps import get_settings
        from app.main import app

        tiny = dataclasses.replace(workflow_settings, max_upload_bytes=16)
        app.dependency_overrides[get_settings] = lambda: tiny
        response = upload(client, plant, small_workbook)
        assert response.status_code == 413
        assert "MDR_MAX_UPLOAD_BYTES" in response.json()["detail"]

    def test_the_service_enforces_the_size_limit_itself(self, plant,
                                                        small_workbook,
                                                        workflow_settings):
        tiny = dataclasses.replace(workflow_settings, max_upload_bytes=16)
        with pytest.raises(workflow_service.UploadTooLarge):
            workflow_service.upload_workbook(
                plant_id=plant.id, filename=small_workbook.name,
                data=small_workbook.read_bytes(), settings=tiny)

    def test_a_missing_form_field_is_a_422(self, client, small_workbook):
        response = client.post(
            "/api/v1/mdr/upload",
            files={"mdr_file": (small_workbook.name, small_workbook.read_bytes())})
        assert response.status_code == 422


# =========================================================== EXTRACT

@pytest.fixture
def uploaded(client, plant, small_workbook) -> dict:
    return upload(client, plant, small_workbook).json()


class TestExtract:

    def test_a_valid_submission_extracts_its_rows(self, client, uploaded):
        response = client.post(f"/api/v1/mdr/{uploaded['mdr_id']}/extract")
        assert response.status_code == 200, response.text
        body = response.json()
        assert body["status"] == SubmissionStatus.EXTRACTED
        assert body["source_sheet_name"] == "QatarEnergy-TN"
        assert body["source_header_row"] == 5
        # Five data rows on the sheet; the spacer row counts as data here
        # because it is not empty. The *document* rows are four - see below.
        assert body["source_row_count"] == 5
        assert body["extracted_at"] is not None

    def test_rows_are_persisted_by_source_row_not_position(
            self, client, uploaded, workflow_settings):
        client.post(f"/api/v1/mdr/{uploaded['mdr_id']}/extract")
        rows = stored_rows(workflow_settings, uploaded["mdr_id"])
        assert [r.source_row for r in rows] == list(DOCUMENT_ROWS)
        # Row 9 is the spacer without a DOCUMENT NO.: no record, no gap-filler.
        assert 9 not in {r.source_row for r in rows}

    def test_the_rows_carry_identity_and_revision(self, client, uploaded,
                                                  workflow_settings,
                                                  small_workbook):
        client.post(f"/api/v1/mdr/{uploaded['mdr_id']}/extract")
        rows = {r.source_row: r
                for r in stored_rows(workflow_settings, uploaded["mdr_id"])}
        engine_docs = {d.source_row: d for d in
                       run_automation(small_workbook).result.documents}
        for source_row, doc in engine_docs.items():
            stored = rows[source_row]
            assert stored.document_identity == doc.document_identity
            assert stored.qatarenergy_document_no == doc.qatarenergy_document_no
            assert stored.revision == doc.revision
            assert stored.revision_raw == doc.revision_raw
            assert stored.is_latest_revision == doc.is_latest_revision
            assert stored.document_title == doc.document_title

    def test_the_automation_columns_are_still_empty_after_extraction(
            self, client, uploaded, workflow_settings):
        """DOC WITH REV, SOW and IDB are the automate step. DOC TYPE is the
        one Phase 2A verdict `MdrEngine.run()` itself produces and the
        repository stores it as the engine handed it over."""
        client.post(f"/api/v1/mdr/{uploaded['mdr_id']}/extract")
        rows = stored_rows(workflow_settings, uploaded["mdr_id"])
        assert {r.doc_with_rev for r in rows} == {""}
        assert {r.sow for r in rows} == {""}
        assert {r.idb_completed_status for r in rows} == {""}
        assert {r.check_status for r in rows} == {""}

    def test_extracting_twice_is_refused_and_duplicates_nothing(
            self, client, uploaded, workflow_settings):
        first = client.post(f"/api/v1/mdr/{uploaded['mdr_id']}/extract")
        second = client.post(f"/api/v1/mdr/{uploaded['mdr_id']}/extract")
        assert first.status_code == 200
        assert second.status_code == 409
        assert "EXTRACTED" in second.json()["detail"]
        assert len(stored_rows(workflow_settings, uploaded["mdr_id"])) == len(
            DOCUMENT_ROWS)

    def test_an_unknown_submission_is_a_404(self, client):
        assert client.post(f"/api/v1/mdr/{uuid.uuid4()}/extract").status_code == 404

    def test_a_workbook_the_reader_cannot_process_fails_and_is_recorded(
            self, client, plant, small_workbook, tmp_path, workflow_settings):
        """Passes the upload probe (the QatarEnergy-TN sheet is there) but
        has no `Status Codes` sheet, which the engine requires. The
        submission ends FAILED with the reader's own reason; nothing is
        stored under it; the uploaded file is kept."""
        wb = load_workbook(small_workbook)
        del wb["Status Codes"]
        broken = tmp_path / "no-status-codes.xlsx"
        wb.save(broken)

        body = upload(client, plant, broken).json()
        response = client.post(f"/api/v1/mdr/{body['mdr_id']}/extract")
        assert response.status_code == 422
        assert "FAILED" in response.json()["detail"]

        stored = stored_submission(workflow_settings, body["mdr_id"])
        assert stored.status == SubmissionStatus.FAILED
        assert "extract" in stored.failure_reason
        assert "Status Codes" in stored.failure_reason
        assert stored_rows(workflow_settings, body["mdr_id"]) == []
        assert (workflow_settings.uploads_dir / stored.stored_path).is_file()

    def test_a_missing_stored_file_is_a_server_failure_and_is_recorded(
            self, client, uploaded, workflow_settings):
        stored = stored_submission(workflow_settings, uploaded["mdr_id"])
        (workflow_settings.uploads_dir / stored.stored_path).unlink()

        response = client.post(f"/api/v1/mdr/{uploaded['mdr_id']}/extract")
        assert response.status_code == 500
        after = stored_submission(workflow_settings, uploaded["mdr_id"])
        assert after.status == SubmissionStatus.FAILED
        assert "missing" in after.failure_reason

    def test_a_failed_submission_cannot_be_extracted_again(
            self, client, uploaded, workflow_settings):
        stored = stored_submission(workflow_settings, uploaded["mdr_id"])
        (workflow_settings.uploads_dir / stored.stored_path).unlink()
        client.post(f"/api/v1/mdr/{uploaded['mdr_id']}/extract")

        response = client.post(f"/api/v1/mdr/{uploaded['mdr_id']}/extract")
        assert response.status_code == 409
        assert "FAILED" in response.json()["detail"]


# ========================================================== AUTOMATE

@pytest.fixture
def extracted(client, uploaded) -> dict:
    assert client.post(f"/api/v1/mdr/{uploaded['mdr_id']}/extract").status_code == 200
    return uploaded


@pytest.fixture(scope="module")
def engine_run(small_workbook):
    """What the existing engine says about the fixture: the oracle."""
    return run_automation(small_workbook)


@requires_rules_workbook
class TestAutomate:

    def test_automation_succeeds_and_reports_the_rule_set(self, client,
                                                          extracted):
        response = client.post(f"/api/v1/mdr/{extracted['mdr_id']}/automate")
        assert response.status_code == 200, response.text
        body = response.json()
        assert body["status"] == SubmissionStatus.AUTOMATED
        assert body["row_count"] == len(DOCUMENT_ROWS)
        assert body["automated_at"] is not None
        assert body["rule_set"]["content_sha256"] == file_digest(RULES_WORKBOOK)
        assert body["rule_set"]["source_filename"] == RULES_WORKBOOK.name

    def test_the_existing_engine_is_what_runs(self, client, extracted,
                                              monkeypatch):
        """`run_automation` - the Phase 2D entry point the CLI uses - is
        called once, with the stored workbook and the rules workbook."""
        calls = []
        real = workflow_service.run_automation

        def spy(workbook, rules_workbook=None):
            calls.append((workbook, rules_workbook))
            return real(workbook, rules_workbook)

        monkeypatch.setattr(workflow_service, "run_automation", spy)
        client.post(f"/api/v1/mdr/{extracted['mdr_id']}/automate")
        assert len(calls) == 1
        workbook, rules = calls[0]
        assert workbook.name == "small-mdr.xlsx"
        assert extracted["mdr_id"] in str(workbook)
        assert rules == RULES_WORKBOOK

    def test_the_four_columns_and_their_provenance_match_the_engine(
            self, client, extracted, workflow_settings, engine_run):
        client.post(f"/api/v1/mdr/{extracted['mdr_id']}/automate")
        stored = {r.source_row: r
                  for r in stored_rows(workflow_settings, extracted["mdr_id"])}
        expected = {r.source_row: r for r in engine_run.rows}
        assert set(stored) == set(expected)
        for source_row, row in expected.items():
            got = stored[source_row]
            assert got.doc_with_rev == row.doc_with_rev
            assert got.doc_type == row.doc_type
            assert got.sow == row.sow
            assert got.idb_completed_status == row.idb_status
            assert got.doc_type_rule == row.doc_type_rule
            assert got.sow_source == row.sow_source
            assert got.idb_source == row.idb_source

    def test_the_columns_are_actually_populated(self, client, extracted,
                                                workflow_settings):
        """The comparison above would pass on all-empty columns if the engine
        produced them; the fixture is chosen so it does not."""
        client.post(f"/api/v1/mdr/{extracted['mdr_id']}/automate")
        rows = stored_rows(workflow_settings, extracted["mdr_id"])
        assert all(r.doc_with_rev for r in rows)
        assert any(r.doc_type for r in rows)
        assert any(r.sow for r in rows)
        assert any(r.idb_completed_status for r in rows)

    def test_check_status_stays_blank(self, client, extracted,
                                      workflow_settings):
        client.post(f"/api/v1/mdr/{extracted['mdr_id']}/automate")
        rows = stored_rows(workflow_settings, extracted["mdr_id"])
        assert {r.check_status for r in rows} == {""}

    def test_a_processing_summary_is_created_and_matches_the_engine(
            self, client, extracted, workflow_settings, engine_run):
        from app.infrastructure.persistence.database import session_scope
        from app.infrastructure.persistence.repositories import (
            ProcessingSummaryRepository,
        )

        client.post(f"/api/v1/mdr/{extracted['mdr_id']}/automate")
        with session_scope(workflow_settings) as session:
            summary = ProcessingSummaryRepository(session).for_submission(
                uuid.UUID(extracted["mdr_id"]))
            rule_set_digest = summary.rule_set.content_sha256
        expected = engine_run.summary()
        assert summary.row_count == expected["rows"]
        assert summary.doc_with_rev_populated == expected["doc_with_rev_populated"]
        assert summary.doc_type_populated == expected["doc_type_populated"]
        assert summary.sow_populated == expected["sow_populated"]
        assert summary.idb_populated == expected["idb_populated"]
        assert summary.idb_manual_check_required == expected[
            "idb_manual_check_required"]
        assert summary.check_status_populated == 0
        assert summary.counts["doc_type_counts"] == expected["doc_type_counts"]
        assert summary.counts["sow_counts"] == expected["sow_counts"]
        assert summary.counts["idb_counts"] == expected["idb_counts"]
        assert summary.counts["discovery"]["qe_sheet"] == "QatarEnergy-TN"
        # Server paths are not persisted, so the summary cannot expose them.
        assert "workbook" not in summary.counts["discovery"]
        assert "rules_workbook" not in summary.counts["discovery"]
        assert summary.engine_version
        assert rule_set_digest == file_digest(RULES_WORKBOOK)

    def test_two_submissions_share_one_rule_set_row(self, client, plant,
                                                    small_workbook):
        ids = []
        for _ in range(2):
            mdr_id = upload(client, plant, small_workbook).json()["mdr_id"]
            client.post(f"/api/v1/mdr/{mdr_id}/extract")
            ids.append(client.post(f"/api/v1/mdr/{mdr_id}/automate").json()
                       ["rule_set"]["rule_set_id"])
        assert ids[0] == ids[1]

    def test_automating_before_extraction_is_refused(self, client, uploaded):
        response = client.post(f"/api/v1/mdr/{uploaded['mdr_id']}/automate")
        assert response.status_code == 409
        assert "not been extracted" in response.json()["detail"]

    def test_automating_twice_is_refused_and_creates_no_second_summary(
            self, client, extracted):
        first = client.post(f"/api/v1/mdr/{extracted['mdr_id']}/automate")
        second = client.post(f"/api/v1/mdr/{extracted['mdr_id']}/automate")
        assert first.status_code == 200
        assert second.status_code == 409
        assert "AUTOMATED" in second.json()["detail"]

    def test_an_unknown_submission_is_a_404(self, client):
        assert client.post(f"/api/v1/mdr/{uuid.uuid4()}/automate").status_code == 404

    def test_an_engine_failure_is_recorded_as_failed_with_the_reason(
            self, client, extracted, workflow_settings, monkeypatch):
        def explode(workbook, rules_workbook=None):
            raise RuntimeError("engine exploded on purpose")

        monkeypatch.setattr(workflow_service, "run_automation", explode)
        response = client.post(f"/api/v1/mdr/{extracted['mdr_id']}/automate")
        assert response.status_code == 500

        stored = stored_submission(workflow_settings, extracted["mdr_id"])
        assert stored.status == SubmissionStatus.FAILED
        assert stored.failure_reason.startswith("automate:")
        assert "engine exploded on purpose" in stored.failure_reason
        # And nothing half-written: the rows still carry no verdicts.
        rows = stored_rows(workflow_settings, extracted["mdr_id"])
        assert {r.doc_with_rev for r in rows} == {""}

    def test_a_row_mismatch_writes_nothing(self, client, extracted,
                                           workflow_settings, monkeypatch):
        """If the engine's rows are not the rows extraction stored, no
        verdict lands on any row and the submission is FAILED."""
        real = workflow_service.run_automation

        def drop_one(workbook, rules_workbook=None):
            run = real(workbook, rules_workbook)
            run.rows.pop()
            return run

        monkeypatch.setattr(workflow_service, "run_automation", drop_one)
        response = client.post(f"/api/v1/mdr/{extracted['mdr_id']}/automate")
        assert response.status_code == 500
        stored = stored_submission(workflow_settings, extracted["mdr_id"])
        assert stored.status == SubmissionStatus.FAILED
        assert "missing by source_row" in stored.failure_reason
        rows = stored_rows(workflow_settings, extracted["mdr_id"])
        assert {r.doc_with_rev for r in rows} == {""}


class TestAutomateWithoutRules:

    def test_no_rules_workbook_is_refused_without_failing_the_submission(
            self, client, extracted, workflow_settings, monkeypatch):
        """A missing rules workbook is a deployment gap, not a fact about
        the submission: 422, and it stays EXTRACTED for when rules arrive."""
        monkeypatch.setattr(type(workflow_settings), "default_rules_workbook",
                            lambda self: None)
        response = client.post(f"/api/v1/mdr/{extracted['mdr_id']}/automate")
        assert response.status_code == 422
        assert "rules workbook" in response.json()["detail"]
        stored = stored_submission(workflow_settings, extracted["mdr_id"])
        assert stored.status == SubmissionStatus.EXTRACTED


# =========================================================== SUMMARY

class TestSummary:

    def test_an_unknown_submission_is_a_404(self, client):
        assert client.get(f"/api/v1/mdr/{uuid.uuid4()}/summary").status_code == 404

    def test_a_malformed_id_is_a_422(self, client):
        assert client.get("/api/v1/mdr/nope/summary").status_code == 422

    def test_an_uploaded_submission_has_a_summary_without_counts(
            self, client, uploaded, plant):
        body = client.get(f"/api/v1/mdr/{uploaded['mdr_id']}/summary").json()
        assert body["status"] == SubmissionStatus.UPLOADED
        assert body["plant_id"] == str(plant.id)
        assert body["plant_code"] == plant.code
        assert body["rule_set"] is None
        assert body["row_count"] == 0
        assert body["counts"] == {}
        assert body["failure_reason"] == ""
        assert body["extracted_at"] is None
        assert body["automated_at"] is None

    @requires_rules_workbook
    def test_an_automated_submission_reports_the_persisted_result(
            self, client, extracted, engine_run):
        automated = client.post(f"/api/v1/mdr/{extracted['mdr_id']}/automate").json()
        body = client.get(f"/api/v1/mdr/{extracted['mdr_id']}/summary").json()
        expected = engine_run.summary()

        assert body["status"] == SubmissionStatus.AUTOMATED
        assert body["row_count"] == expected["rows"] == len(DOCUMENT_ROWS)
        assert body["doc_type_populated"] == expected["doc_type_populated"]
        assert body["sow_populated"] == expected["sow_populated"]
        assert body["sow_unresolved"] == expected["sow_unresolved"]
        assert body["idb_populated"] == expected["idb_populated"]
        assert body["idb_unmapped"] == expected["idb_unmapped"]
        assert body["idb_manual_check_required"] == expected[
            "idb_manual_check_required"]
        assert body["check_status_populated"] == 0
        assert body["counts"]["doc_type_counts"] == expected["doc_type_counts"]
        assert body["counts"]["sow_counts"] == expected["sow_counts"]
        assert body["counts"]["idb_counts"] == expected["idb_counts"]
        assert body["rule_set"] == automated["rule_set"]
        assert body["engine_version"]
        assert body["extracted_at"] and body["automated_at"]

    def test_a_failed_submission_reports_its_reason(self, client, uploaded,
                                                    workflow_settings):
        stored = stored_submission(workflow_settings, uploaded["mdr_id"])
        (workflow_settings.uploads_dir / stored.stored_path).unlink()
        client.post(f"/api/v1/mdr/{uploaded['mdr_id']}/extract")

        body = client.get(f"/api/v1/mdr/{uploaded['mdr_id']}/summary").json()
        assert body["status"] == SubmissionStatus.FAILED
        assert body["failure_reason"].startswith("extract:")
        assert body["rule_set"] is None

    @requires_rules_workbook
    def test_the_summary_runs_nothing_and_changes_nothing(
            self, client, extracted, workflow_settings, monkeypatch):
        client.post(f"/api/v1/mdr/{extracted['mdr_id']}/automate")
        before = stored_submission(workflow_settings, extracted["mdr_id"])

        def forbidden(*args, **kwargs):
            raise AssertionError("summary must not run the engine")

        monkeypatch.setattr(workflow_service, "run_automation", forbidden)
        monkeypatch.setattr(workflow_service, "MdrEngine", forbidden)
        first = client.get(f"/api/v1/mdr/{extracted['mdr_id']}/summary")
        second = client.get(f"/api/v1/mdr/{extracted['mdr_id']}/summary")
        assert first.status_code == second.status_code == 200
        assert first.json() == second.json()

        after = stored_submission(workflow_settings, extracted["mdr_id"])
        assert after.status == before.status
        assert after.automated_at == before.automated_at


# ======================================================== END TO END

@requires_rules_workbook
class TestEndToEnd:

    def test_upload_extract_automate_summary(self, client, plant,
                                             small_workbook, workflow_settings,
                                             engine_run):
        """The whole workflow, once, and the database against the engine."""
        source_before = hashlib.sha256(small_workbook.read_bytes()).hexdigest()

        up = upload(client, plant, small_workbook)
        assert up.status_code == 201
        mdr_id = up.json()["mdr_id"]

        ex = client.post(f"/api/v1/mdr/{mdr_id}/extract")
        assert ex.status_code == 200 and ex.json()["status"] == "EXTRACTED"

        au = client.post(f"/api/v1/mdr/{mdr_id}/automate")
        assert au.status_code == 200 and au.json()["status"] == "AUTOMATED"

        su = client.get(f"/api/v1/mdr/{mdr_id}/summary")
        assert su.status_code == 200
        assert su.json()["row_count"] == len(engine_run.rows)

        stored = {r.source_row: r for r in stored_rows(workflow_settings, mdr_id)}
        for row in engine_run.rows:
            assert stored[row.source_row].doc_with_rev == row.doc_with_rev
            assert stored[row.source_row].doc_type == row.doc_type
            assert stored[row.source_row].sow == row.sow
            assert stored[row.source_row].idb_completed_status == row.idb_status
            assert stored[row.source_row].check_status == ""

        assert hashlib.sha256(
            small_workbook.read_bytes()).hexdigest() == source_before
        assert client.get(f"/api/v1/mdr/{mdr_id}/download").status_code == 501

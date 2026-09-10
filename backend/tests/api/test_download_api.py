"""`GET /api/v1/mdr/{mdr_id}/download` over HTTP, on the small fixture.

Delivery Phase 4. Every test goes through the real route, the real service,
the real Phase 1 writer and the real PostgreSQL schema; only the workbook is
small. The generated file is compared with what the database holds, because
the whole point of the endpoint is that the downloaded sheet *is* the
persisted result - not a re-run, and not a reconstruction of one.

The real ~22k-row workbook goes through the same path in
`test_workflow_real_workbook.py`, where the Phase 1 counts are pinned.
"""

from __future__ import annotations

import hashlib
import io
import re
import tempfile
import urllib.parse
import uuid
from pathlib import Path

import pytest
from openpyxl import load_workbook
from sqlalchemy import text

from app.domain.enums.lifecycle import SubmissionStatus
from app.domain.models.automation import (
    AUTOMATION_COLUMNS, CHECK_STATUS, DOC_IDB_COMPLETED_STATUS,
    DOC_IS_REQUIRED_SOW, DOC_TYPE, DOC_WITH_REV,
)
from app.infrastructure.excel.output_workbook import (
    AUTOMATED_SHEET, read_automation_columns,
)
from app.infrastructure.persistence.database import session_scope
from app.services import workflow_service
from app.services.export_service import AUTOMATED_WORKBOOK_SUFFIX
from tests.support.automation import (
    DATA_ROWS, DOCUMENT_ROWS, FIRST_DATA_ROW, HEADER_ROW, HEADERS,
)
from tests.support.workbook import requires_rules_workbook

from .conftest import stored_rows, stored_submission, upload

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
SOURCE_SHEETS = ["QatarEnergy-TN", "TN FROM VENDORS", "Status Codes",
                 "VENDOR LIST"]
#: Where the five columns land on the fixture - see
#: `tests/integration/test_automated_workbook.py`.
AUTOMATION_LETTERS = ["I", "J", "K", "L", "M"]

pytestmark = requires_rules_workbook


# ---------------------------------------------------------------- fixtures

@pytest.fixture
def uploaded(client, plant, small_workbook) -> dict:
    return upload(client, plant, small_workbook).json()


@pytest.fixture
def extracted(client, uploaded) -> dict:
    assert client.post(f"/api/v1/mdr/{uploaded['mdr_id']}/extract").status_code == 200
    return uploaded


@pytest.fixture
def automated(client, extracted) -> dict:
    response = client.post(f"/api/v1/mdr/{extracted['mdr_id']}/automate")
    assert response.status_code == 200, response.text
    return extracted


@pytest.fixture
def downloaded(client, automated, tmp_path) -> tuple[dict, object, Path]:
    """One download of an automated submission: (upload body, response,
    the body saved to disk so openpyxl can open it)."""
    response = client.get(f"/api/v1/mdr/{automated['mdr_id']}/download")
    assert response.status_code == 200, response.text
    path = tmp_path / "downloaded.xlsx"
    path.write_bytes(response.content)
    return automated, response, path


def _disposition_filename(response) -> str:
    """The filename a client would save under, from either RFC 6266 form."""
    header = response.headers["content-disposition"]
    assert header.startswith("attachment;")
    extended = re.search(r"filename\*=utf-8''([^;]+)", header, re.I)
    if extended:
        return urllib.parse.unquote(extended.group(1))
    plain = re.search(r'filename="([^"]*)"', header)
    assert plain, header
    return plain.group(1)


def _sql(settings, statement: str, **params) -> None:
    """One statement against the throwaway database, committed. Used only to
    manufacture the inconsistent states the endpoint must refuse."""
    with session_scope(settings) as session:
        session.execute(text(statement), params)


# ================================================================ SUCCESS

class TestASuccessfulDownload:

    def test_an_automated_submission_downloads(self, downloaded):
        _, response, _ = downloaded
        assert response.status_code == 200
        assert len(response.content) > 0

    def test_the_content_type_is_xlsx(self, downloaded):
        _, response, _ = downloaded
        assert response.headers["content-type"].split(";")[0] == XLSX
        assert response.headers["content-length"] == str(len(response.content))

    def test_the_body_is_a_zip_container_like_every_xlsx(self, downloaded):
        _, response, _ = downloaded
        assert response.content[:2] == b"PK"

    def test_the_filename_is_the_uploads_name_plus_the_export_suffix(
            self, downloaded, small_workbook):
        """The same name the CLI's `--excel` would give the file."""
        _, response, _ = downloaded
        assert _disposition_filename(response) == \
            small_workbook.stem + AUTOMATED_WORKBOOK_SUFFIX
        assert _disposition_filename(response).endswith(".xlsx")

    def test_the_filename_carries_no_path_and_no_header_breaking_character(
            self, downloaded, workflow_settings):
        _, response, _ = downloaded
        header = response.headers["content-disposition"]
        name = _disposition_filename(response)
        for forbidden in ("/", "\\", '"', ";", "\r", "\n", ".."):
            assert forbidden not in name, name
        assert str(workflow_settings.uploads_dir) not in header
        assert tempfile.gettempdir() not in header

    def test_a_hostile_upload_name_becomes_a_safe_download_name(
            self, client, plant, small_workbook):
        """A client may send a path or header separators as the filename.
        The download name keeps the final component and only the characters
        `safe_filename` allows."""
        body = upload(client, plant, small_workbook,
                      filename="..\\..\\evil name;drop.xlsx").json()
        client.post(f"/api/v1/mdr/{body['mdr_id']}/extract")
        assert client.post(f"/api/v1/mdr/{body['mdr_id']}/automate").status_code == 200
        response = client.get(f"/api/v1/mdr/{body['mdr_id']}/download")
        assert response.status_code == 200
        name = _disposition_filename(response)
        assert name == "evil name_drop" + AUTOMATED_WORKBOOK_SUFFIX
        assert re.fullmatch(r"[A-Za-z0-9 ._()&#\-]+", name), name

    def test_the_body_opens_in_openpyxl(self, downloaded):
        _, response, _ = downloaded
        wb = load_workbook(io.BytesIO(response.content))
        try:
            assert wb.sheetnames
        finally:
            wb.close()

    def test_downloading_twice_gives_the_same_sheet(self, client, downloaded):
        """Ephemeral generation: nothing is cached and nothing is stored, so
        the second download is generated again - and is the same result."""
        body, _, first = downloaded
        again = client.get(f"/api/v1/mdr/{body['mdr_id']}/download")
        assert again.status_code == 200
        second = first.parent / "again.xlsx"
        second.write_bytes(again.content)
        assert read_automation_columns(first) == read_automation_columns(second)


# ====================================================== THE WORKBOOK ITSELF

class TestTheOriginalWorkbookIsPreserved:

    def test_every_original_sheet_is_there_in_order_plus_the_automated_one(
            self, downloaded):
        _, _, path = downloaded
        wb = load_workbook(path, read_only=True)
        try:
            assert wb.sheetnames == SOURCE_SHEETS + [AUTOMATED_SHEET]
        finally:
            wb.close()

    def test_the_original_sheets_are_cell_for_cell_unchanged(
            self, downloaded, small_workbook):
        _, _, path = downloaded

        def cells(file, sheet):
            wb = load_workbook(file, data_only=True)
            try:
                return [[c.value for c in row] for row in wb[sheet].rows]
            finally:
                wb.close()

        for sheet in SOURCE_SHEETS:
            assert cells(path, sheet) == cells(small_workbook, sheet), sheet

    def test_the_source_sheet_did_not_gain_the_automation_columns(
            self, downloaded):
        _, _, path = downloaded
        wb = load_workbook(path, read_only=True)
        try:
            captions = [c.value for c in
                        next(wb["QatarEnergy-TN"].iter_rows(
                            min_row=HEADER_ROW, max_row=HEADER_ROW))]
        finally:
            wb.close()
        assert [c for c in captions if c] == list(HEADERS)
        assert not set(AUTOMATION_COLUMNS) & set(captions)

    def test_the_formula_survives_as_a_formula(self, downloaded):
        _, _, path = downloaded
        wb = load_workbook(path)                    # not data_only
        try:
            assert wb["QatarEnergy-TN"]["K4"].value == "=NOW()"
            # Five columns were inserted to its left on the automated sheet.
            assert wb[AUTOMATED_SHEET]["P4"].value == "=NOW()"
        finally:
            wb.close()

    def test_formatting_merges_filter_pane_cf_and_dv_survive(self, downloaded):
        """The Phase 1 writer's guarantees, observed on the downloaded file:
        the header style, the merged title band (and the merge straddling
        the insertion, widened), the frozen pane, the auto-filter widened
        over the new columns, and the conditional format and validation on
        the trailing block moved right with it."""
        _, _, path = downloaded
        wb = load_workbook(path)
        try:
            ws = wb[AUTOMATED_SHEET]
            assert ws["A5"].font.bold is True
            assert ws["I5"].font.bold is True                # new caption
            assert ws["I5"].fill.fgColor.rgb == "FF00B050"
            ranges = {str(r) for r in ws.merged_cells.ranges}
            assert {"A1:E1", "A3:C3", "H4:N4"} <= ranges
            assert ws.freeze_panes == "B6"
            assert ws.auto_filter.ref == "A5:P10"
            assert "O6:O10" in {str(cf.sqref) for cf in ws.conditional_formatting}
            assert "N6:N10" in {str(dv.sqref) for dv in
                                ws.data_validations.dataValidation}
            assert ws.column_dimensions["B"].width == 30.0
            assert ws.column_dimensions["N"].width == 26.0    # was I

            src = wb["QatarEnergy-TN"]
            assert src.freeze_panes == "B6"
            assert src.auto_filter.ref == "A5:K10"
            assert {"A1:E1", "A3:C3", "H4:I4"} <= {
                str(r) for r in src.merged_cells.ranges}
        finally:
            wb.close()

    def test_the_row_count_is_preserved(self, downloaded):
        _, _, path = downloaded
        wb = load_workbook(path, data_only=True)
        try:
            expected = FIRST_DATA_ROW + len(DATA_ROWS) - 1
            assert wb["QatarEnergy-TN"].max_row == expected
            assert wb[AUTOMATED_SHEET].max_row == expected
        finally:
            wb.close()


class TestTheAutomatedSheet:

    def test_it_exists_under_the_contract_name(self, downloaded):
        _, _, path = downloaded
        wb = load_workbook(path, read_only=True)
        try:
            assert "QatarEnergy-TN Automated" in wb.sheetnames
            assert wb.sheetnames[-1] == "QatarEnergy-TN Automated"
        finally:
            wb.close()

    def test_the_five_captions_are_at_the_working_positions(self, downloaded):
        _, _, path = downloaded
        _, letters, _ = read_automation_columns(path)
        assert [letters[c] for c in AUTOMATION_COLUMNS] == AUTOMATION_LETTERS

    def test_the_values_are_the_persisted_ones_row_for_row(
            self, downloaded, workflow_settings):
        """Every cell of the four columns against `mdr_document_rows`, keyed
        by `source_row`. Not against the engine: the database is the
        authority the download serves from."""
        body, _, path = downloaded
        stored = {r.source_row: r
                  for r in stored_rows(workflow_settings, body["mdr_id"])}
        _, _, rows = read_automation_columns(path)
        by_row = {r["source_row"]: r for r in rows}
        assert set(by_row) == set(stored) == set(DOCUMENT_ROWS)
        for source_row, row in stored.items():
            got = by_row[source_row]
            assert got[DOC_WITH_REV] == row.doc_with_rev.upper()
            assert got[DOC_TYPE] == row.doc_type.upper()
            assert got[DOC_IS_REQUIRED_SOW] == row.sow.upper()
            assert got[DOC_IDB_COMPLETED_STATUS] == \
                row.idb_completed_status.upper()
            assert got[CHECK_STATUS] == ""

    def test_the_values_are_populated_not_merely_present(self, downloaded):
        _, _, path = downloaded
        _, _, rows = read_automation_columns(path)
        assert all(r[DOC_WITH_REV] for r in rows)
        assert any(r[DOC_TYPE] for r in rows)
        assert any(r[DOC_IS_REQUIRED_SOW] for r in rows)
        assert any(r[DOC_IDB_COMPLETED_STATUS] for r in rows)

    def test_the_skipped_row_keeps_five_empty_cells(self, downloaded):
        """Row 9 has no DOCUMENT NO. and no persisted row. Mapping is by
        `source_row`, so its neighbours' answers do not shift onto it."""
        _, _, path = downloaded
        wb = load_workbook(path, data_only=True)
        try:
            ws = wb[AUTOMATED_SHEET]
            assert 9 not in DOCUMENT_ROWS
            assert [ws[f"{L}9"].value for L in AUTOMATION_LETTERS] == [None] * 5
            # And the rows around it carry their own answers.
            assert ws["I8"].value == "4391-MG-WPR-0041-1"
            assert ws["I10"].value == "VEN-MEWTP-5-43-0011"
        finally:
            wb.close()

    def test_every_document_row_and_only_those_carry_values(self, downloaded):
        _, _, path = downloaded
        wb = load_workbook(path, data_only=True)
        try:
            ws = wb[AUTOMATED_SHEET]
            populated = [r for r in range(HEADER_ROW + 1, ws.max_row + 1)
                         if ws[f"I{r}"].value]
        finally:
            wb.close()
        assert populated == list(DOCUMENT_ROWS)

    def test_check_status_is_blank_on_every_row(self, downloaded):
        _, _, path = downloaded
        wb = load_workbook(path, data_only=True)
        try:
            ws = wb[AUTOMATED_SHEET]
            assert ws["M5"].value == CHECK_STATUS
            assert [ws[f"M{r}"].value
                    for r in range(HEADER_ROW + 1, ws.max_row + 1)] == \
                [None] * len(DATA_ROWS)
        finally:
            wb.close()


# ====================================================== NOTHING IS CHANGED

class TestNothingIsChangedByDownloading:

    def test_the_stored_upload_is_not_modified(self, client, automated,
                                               workflow_settings):
        stored = stored_submission(workflow_settings, automated["mdr_id"])
        path = workflow_settings.uploads_dir / stored.stored_path
        before = (hashlib.sha256(path.read_bytes()).hexdigest(),
                  path.stat().st_size, path.stat().st_mtime_ns)
        assert client.get(f"/api/v1/mdr/{automated['mdr_id']}/download").status_code == 200
        after = (hashlib.sha256(path.read_bytes()).hexdigest(),
                 path.stat().st_size, path.stat().st_mtime_ns)
        assert after == before
        assert before[0] == stored.source_sha256

    def test_the_stored_upload_still_has_no_automated_sheet(
            self, downloaded, workflow_settings):
        body, _, _ = downloaded
        stored = stored_submission(workflow_settings, body["mdr_id"])
        wb = load_workbook(workflow_settings.uploads_dir / stored.stored_path,
                           read_only=True)
        try:
            assert wb.sheetnames == SOURCE_SHEETS
        finally:
            wb.close()

    def test_the_fixture_file_itself_is_untouched(self, client, plant,
                                                  small_workbook):
        digest = hashlib.sha256(small_workbook.read_bytes()).hexdigest()
        body = upload(client, plant, small_workbook).json()
        client.post(f"/api/v1/mdr/{body['mdr_id']}/extract")
        client.post(f"/api/v1/mdr/{body['mdr_id']}/automate")
        client.get(f"/api/v1/mdr/{body['mdr_id']}/download")
        assert hashlib.sha256(small_workbook.read_bytes()).hexdigest() == digest

    def test_no_engine_runs(self, client, automated, monkeypatch):
        """The download rebuilds the sheet from the database. If it reached
        for the engine the request would fail here."""
        def forbidden(*args, **kwargs):
            raise AssertionError("download must not run the engine")

        monkeypatch.setattr(workflow_service, "run_automation", forbidden)
        monkeypatch.setattr(workflow_service, "MdrEngine", forbidden)
        monkeypatch.setattr(workflow_service, "register_rule_set", forbidden)
        response = client.get(f"/api/v1/mdr/{automated['mdr_id']}/download")
        assert response.status_code == 200

    def test_the_writer_receives_the_persisted_rows_not_engine_rows(
            self, client, automated, workflow_settings, monkeypatch):
        """Belt and braces: what is handed to the Phase 1 writer is what the
        database holds, row for row, provenance included."""
        handed = []
        real = workflow_service.export_automated_workbook

        def spy(source, rows, outdir, destination=None):
            rows = list(rows)
            handed.append(rows)
            return real(source, rows, outdir, destination)

        monkeypatch.setattr(workflow_service, "export_automated_workbook", spy)
        assert client.get(f"/api/v1/mdr/{automated['mdr_id']}/download").status_code == 200
        assert len(handed) == 1
        stored = {r.source_row: r
                  for r in stored_rows(workflow_settings, automated["mdr_id"])}
        assert {r.source_row for r in handed[0]} == set(stored)
        for row in handed[0]:
            s = stored[row.source_row]
            assert (row.doc_with_rev, row.doc_type, row.sow, row.idb_status,
                    row.doc_type_rule, row.sow_source, row.idb_source) == (
                s.doc_with_rev, s.doc_type, s.sow, s.idb_completed_status,
                s.doc_type_rule, s.sow_source, s.idb_source)
            assert row.check_status == ""

    def test_the_database_is_unchanged(self, client, automated,
                                       workflow_settings):
        before = stored_submission(workflow_settings, automated["mdr_id"])
        rows_before = [(r.id, r.source_row, r.doc_with_rev, r.doc_type, r.sow,
                        r.idb_completed_status, r.check_status)
                       for r in stored_rows(workflow_settings, automated["mdr_id"])]
        summary_before = client.get(f"/api/v1/mdr/{automated['mdr_id']}/summary").json()

        assert client.get(f"/api/v1/mdr/{automated['mdr_id']}/download").status_code == 200

        after = stored_submission(workflow_settings, automated["mdr_id"])
        assert after.status == before.status == SubmissionStatus.AUTOMATED
        assert after.automated_at == before.automated_at
        assert after.updated_at == before.updated_at
        assert after.stored_path == before.stored_path
        assert [(r.id, r.source_row, r.doc_with_rev, r.doc_type, r.sow,
                 r.idb_completed_status, r.check_status)
                for r in stored_rows(workflow_settings, automated["mdr_id"])] == rows_before
        assert client.get(f"/api/v1/mdr/{automated['mdr_id']}/summary").json() == summary_before

    def test_nothing_is_written_under_the_uploads_directory(
            self, client, automated, workflow_settings):
        root = workflow_settings.uploads_dir
        before = sorted(p for p in root.rglob("*") if p.is_file())
        assert client.get(f"/api/v1/mdr/{automated['mdr_id']}/download").status_code == 200
        assert sorted(p for p in root.rglob("*") if p.is_file()) == before
        assert not list(root.rglob(f"*{AUTOMATED_WORKBOOK_SUFFIX}"))


# ================================================================ REFUSALS

class TestRefusals:

    def test_an_unknown_submission_is_a_404(self, client):
        response = client.get(f"/api/v1/mdr/{uuid.uuid4()}/download")
        assert response.status_code == 404
        assert set(response.json()) == {"detail"}

    def test_a_malformed_id_is_a_422(self, client):
        assert client.get("/api/v1/mdr/not-a-uuid/download").status_code == 422

    def test_an_uploaded_submission_is_refused(self, client, uploaded):
        response = client.get(f"/api/v1/mdr/{uploaded['mdr_id']}/download")
        assert response.status_code == 409
        assert "not been extracted" in response.json()["detail"]
        assert "AUTOMATED" in response.json()["detail"]

    def test_an_extracted_submission_is_refused(self, client, extracted):
        response = client.get(f"/api/v1/mdr/{extracted['mdr_id']}/download")
        assert response.status_code == 409
        assert "not been automated" in response.json()["detail"]

    def test_a_failed_submission_is_refused(self, client, uploaded,
                                            workflow_settings):
        stored = stored_submission(workflow_settings, uploaded["mdr_id"])
        (workflow_settings.uploads_dir / stored.stored_path).unlink()
        assert client.post(f"/api/v1/mdr/{uploaded['mdr_id']}/extract").status_code == 500
        assert stored_submission(workflow_settings,
                                 uploaded["mdr_id"]).status == SubmissionStatus.FAILED

        response = client.get(f"/api/v1/mdr/{uploaded['mdr_id']}/download")
        assert response.status_code == 409
        assert "FAILED" in response.json()["detail"]

    def test_a_refusal_carries_no_body_but_the_detail(self, client, extracted):
        body = client.get(f"/api/v1/mdr/{extracted['mdr_id']}/download").json()
        assert set(body) == {"detail"}


# ======================================================= SERVER FAILURES

class TestServerFailures:
    """Each is a 500 whose detail names the submission and never a server
    path, and each leaves the submission AUTOMATED: download is a read."""

    @staticmethod
    def _still_automated(settings, mdr_id) -> None:
        after = stored_submission(settings, mdr_id)
        assert after.status == SubmissionStatus.AUTOMATED
        assert after.failure_reason == ""

    @staticmethod
    def _exposes_no_path(response, settings) -> None:
        detail = response.json()["detail"]
        assert str(settings.uploads_dir) not in detail
        assert tempfile.gettempdir() not in detail
        assert "\\" not in detail and "/tmp" not in detail

    def test_a_missing_stored_workbook_is_a_500(self, client, automated,
                                                workflow_settings):
        stored = stored_submission(workflow_settings, automated["mdr_id"])
        (workflow_settings.uploads_dir / stored.stored_path).unlink()

        response = client.get(f"/api/v1/mdr/{automated['mdr_id']}/download")
        assert response.status_code == 500
        assert "missing" in response.json()["detail"]
        self._exposes_no_path(response, workflow_settings)
        self._still_automated(workflow_settings, automated["mdr_id"])

    def test_a_corrupt_stored_workbook_is_a_500(self, client, automated,
                                                workflow_settings):
        """Truncated after upload. Caught by the digest check before the
        writer opens it, and reported as what it is."""
        stored = stored_submission(workflow_settings, automated["mdr_id"])
        path = workflow_settings.uploads_dir / stored.stored_path
        path.write_bytes(path.read_bytes()[:1000])

        response = client.get(f"/api/v1/mdr/{automated['mdr_id']}/download")
        assert response.status_code == 500
        assert "digest" in response.json()["detail"]
        self._exposes_no_path(response, workflow_settings)
        self._still_automated(workflow_settings, automated["mdr_id"])

    def test_an_unreadable_workbook_the_digest_check_lets_through_is_a_500(
            self, client, automated, workflow_settings, monkeypatch):
        """If the bytes somehow pass the digest and still cannot be opened
        for writing, the writer's failure is reported without its message."""
        stored = stored_submission(workflow_settings, automated["mdr_id"])
        monkeypatch.setattr(workflow_service, "sha256_file",
                            lambda path: stored.source_sha256)
        path = workflow_settings.uploads_dir / stored.stored_path
        path.write_bytes(b"PK\x03\x04 not a workbook any more")

        response = client.get(f"/api/v1/mdr/{automated['mdr_id']}/download")
        assert response.status_code == 500
        assert "could not be generated" in response.json()["detail"]
        self._exposes_no_path(response, workflow_settings)
        self._still_automated(workflow_settings, automated["mdr_id"])

    def test_a_writer_failure_is_a_500_and_leaves_no_temporary_file(
            self, client, automated, workflow_settings, monkeypatch):
        def explode(source, rows, outdir, destination=None):
            raise RuntimeError(f"writer exploded on purpose at {outdir}")

        monkeypatch.setattr(workflow_service, "export_automated_workbook", explode)
        tmp = Path(tempfile.gettempdir())
        before = set(tmp.glob("mdr-download-*"))

        response = client.get(f"/api/v1/mdr/{automated['mdr_id']}/download")
        assert response.status_code == 500
        assert "RuntimeError" in response.json()["detail"]
        assert "exploded" not in response.json()["detail"]     # message withheld
        self._exposes_no_path(response, workflow_settings)
        assert set(tmp.glob("mdr-download-*")) == before
        self._still_automated(workflow_settings, automated["mdr_id"])


class TestInconsistentPersistedRows:
    """A stored result that does not add up is refused, never patched."""

    def test_a_row_missing_from_the_database_is_a_500(self, client, automated,
                                                      workflow_settings):
        _sql(workflow_settings,
             "DELETE FROM mdr_document_rows WHERE submission_id = :id "
             "AND source_row = :row", id=automated["mdr_id"], row=DOCUMENT_ROWS[-1])
        response = client.get(f"/api/v1/mdr/{automated['mdr_id']}/download")
        assert response.status_code == 500
        detail = response.json()["detail"]
        assert f"{len(DOCUMENT_ROWS) - 1} document rows" in detail
        assert f"records {len(DOCUMENT_ROWS)}" in detail

    def test_no_rows_at_all_is_a_500(self, client, automated, workflow_settings):
        _sql(workflow_settings,
             "DELETE FROM mdr_document_rows WHERE submission_id = :id",
             id=automated["mdr_id"])
        response = client.get(f"/api/v1/mdr/{automated['mdr_id']}/download")
        assert response.status_code == 500
        assert "no document rows" in response.json()["detail"]

    def test_a_row_whose_verdict_disagrees_with_the_summary_is_a_500(
            self, client, automated, workflow_settings):
        """One row has lost its DOC WITH REV: the summary says every row has
        one. The download refuses rather than serving a sheet that
        silently disagrees with the summary the client was shown."""
        _sql(workflow_settings,
             "UPDATE mdr_document_rows SET doc_with_rev = '' "
             "WHERE submission_id = :id AND source_row = :row",
             id=automated["mdr_id"], row=DOCUMENT_ROWS[0])
        response = client.get(f"/api/v1/mdr/{automated['mdr_id']}/download")
        assert response.status_code == 500
        assert "doc_with_rev_populated" in response.json()["detail"]

    def test_a_row_at_or_above_the_header_row_is_a_500(self, client, automated,
                                                       workflow_settings):
        _sql(workflow_settings,
             "UPDATE mdr_document_rows SET source_row = :bad "
             "WHERE submission_id = :id AND source_row = :row",
             id=automated["mdr_id"], row=DOCUMENT_ROWS[0], bad=HEADER_ROW)
        response = client.get(f"/api/v1/mdr/{automated['mdr_id']}/download")
        assert response.status_code == 500
        assert "header row" in response.json()["detail"]

    def test_a_missing_header_row_is_a_500(self, client, automated,
                                           workflow_settings):
        _sql(workflow_settings,
             "UPDATE mdr_submissions SET source_header_row = NULL WHERE id = :id",
             id=automated["mdr_id"])
        response = client.get(f"/api/v1/mdr/{automated['mdr_id']}/download")
        assert response.status_code == 500
        assert "header row" in response.json()["detail"]

    def test_a_header_row_the_workbook_does_not_have_is_a_500(
            self, client, automated, workflow_settings):
        """The recorded header row is not where the stored workbook's header
        is: the rows were not extracted from this file."""
        _sql(workflow_settings,
             "UPDATE mdr_submissions SET source_header_row = :row WHERE id = :id",
             id=automated["mdr_id"], row=HEADER_ROW - 2)
        response = client.get(f"/api/v1/mdr/{automated['mdr_id']}/download")
        assert response.status_code == 500
        assert "extraction recorded row" in response.json()["detail"]

    def test_a_missing_summary_is_a_500(self, client, automated,
                                        workflow_settings):
        _sql(workflow_settings,
             "DELETE FROM mdr_processing_summaries WHERE submission_id = :id",
             id=automated["mdr_id"])
        response = client.get(f"/api/v1/mdr/{automated['mdr_id']}/download")
        assert response.status_code == 500
        assert "no processing summary" in response.json()["detail"]

    def test_duplicate_source_rows_are_a_500(self, client, automated,
                                             monkeypatch):
        """The unique constraint makes this impossible in PostgreSQL; the
        service still refuses rather than letting the writer pick one."""
        from app.infrastructure.persistence.repositories import (
            DocumentRowRepository,
        )
        real = DocumentRowRepository.for_submission

        def duplicated(self, submission_id, **kwargs):
            rows = list(real(self, submission_id, **kwargs))
            return rows + rows[:1]

        monkeypatch.setattr(DocumentRowRepository, "for_submission", duplicated)
        response = client.get(f"/api/v1/mdr/{automated['mdr_id']}/download")
        assert response.status_code == 500
        assert "duplicate" in response.json()["detail"]

    def test_the_inconsistency_does_not_fail_the_submission(
            self, client, automated, workflow_settings):
        _sql(workflow_settings,
             "DELETE FROM mdr_document_rows WHERE submission_id = :id "
             "AND source_row = :row", id=automated["mdr_id"], row=DOCUMENT_ROWS[-1])
        client.get(f"/api/v1/mdr/{automated['mdr_id']}/download")
        after = stored_submission(workflow_settings, automated["mdr_id"])
        assert after.status == SubmissionStatus.AUTOMATED
        assert after.failure_reason == ""


# ========================================================= TEMPORARY FILE

class TestTheTemporaryFileIsCleanedUp:

    def test_the_generated_file_is_gone_once_the_response_is_sent(
            self, client, automated, monkeypatch):
        artifacts = []
        real = workflow_service.download_submission

        def spy(mdr_id, **kwargs):
            artifact = real(mdr_id, **kwargs)
            artifacts.append(artifact)
            return artifact

        monkeypatch.setattr(workflow_service, "download_submission", spy)
        response = client.get(f"/api/v1/mdr/{automated['mdr_id']}/download")
        assert response.status_code == 200
        assert len(artifacts) == 1
        artifact = artifacts[0]
        assert artifact.workdir.name.startswith("mdr-download-")
        assert not artifact.path.exists()
        assert not artifact.workdir.exists()

    def test_the_file_was_generated_outside_every_data_directory(
            self, client, automated, workflow_settings, monkeypatch):
        seen = []
        real = workflow_service.download_submission

        def spy(mdr_id, **kwargs):
            artifact = real(mdr_id, **kwargs)
            seen.append((artifact.path, artifact.path.is_file()))
            return artifact

        monkeypatch.setattr(workflow_service, "download_submission", spy)
        client.get(f"/api/v1/mdr/{automated['mdr_id']}/download")
        path, existed = seen[0]
        assert existed
        for root in (workflow_settings.uploads_dir, workflow_settings.data_dir):
            assert root.resolve() not in path.resolve().parents
        assert Path(tempfile.gettempdir()).resolve() in path.resolve().parents

    def test_no_temporary_directory_is_left_behind_by_a_success(
            self, client, automated):
        tmp = Path(tempfile.gettempdir())
        before = set(tmp.glob("mdr-download-*"))
        assert client.get(f"/api/v1/mdr/{automated['mdr_id']}/download").status_code == 200
        assert set(tmp.glob("mdr-download-*")) == before

    def test_cleanup_is_idempotent(self, automated, workflow_settings):
        artifact = workflow_service.download_submission(
            uuid.UUID(automated["mdr_id"]), settings=workflow_settings)
        assert artifact.path.is_file()
        artifact.cleanup()
        assert not artifact.workdir.exists()
        artifact.cleanup()                                  # nothing to do
        assert not artifact.workdir.exists()

    def test_cleanup_runs_even_when_the_send_fails(self, tmp_path):
        """A client that disconnects mid-stream makes the ASGI send raise.
        The response still removes its file."""
        import asyncio

        from app.api.routes.mdr_v1 import _EphemeralFileResponse

        path = tmp_path / "body.xlsx"
        path.write_bytes(b"PK\x03\x04")
        cleaned = []
        response = _EphemeralFileResponse(
            path, media_type=XLSX, filename="body.xlsx",
            cleanup=lambda: cleaned.append(True))

        async def send(message):
            raise ConnectionError("client went away")

        async def receive():                       # pragma: no cover
            return {"type": "http.request"}

        scope = {"type": "http", "method": "GET", "headers": []}
        with pytest.raises(ConnectionError):
            asyncio.run(response(scope, receive, send))
        assert cleaned == [True]

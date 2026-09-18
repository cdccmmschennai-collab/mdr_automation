"""A small MDR workbook with multiple revisions per document.

`tests/support/automation.py`'s fixture is deliberately one row per document -
it exists to pin the Phase 2D writer's formatting. This one exists to pin
*revision grouping*: several rows share a `DOCUMENT NO.` and only the highest
`(band, ordinal)` - the same key `engine.revision.ranking.determine_latest`
already ranks by - should end up in `Latest Revisions`.

Kept deliberately plain (a single header row, no title band, no merges): the
formatting guarantees are already covered by `tests/support/automation.py`;
what this fixture is for is revision data.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook

HEADERS: tuple[str, ...] = (
    "DOCUMENT NO.", "REV", "DISCIPLINE", "DOCUMENT TITLE",
    "ISSUED STATUS IN PDMS", "LATEST/ NOT LATEST", "REMARKS",
    "QATARENERGY SIGNED / NOT SIGNED",
)

HEADER_ROW = 1
FIRST_DATA_ROW = 2

#: (document_no, rev) rows, in sheet order, for the mixed document set the
#: focused tests share. Excel row numbers follow from position (row 2..).
#:
#: DOC-NUM-001: a numeric progression, 0 -> 1 -> 2.
#: DOC-ALP-002: an alphabetic progression, A -> B.
#: DOC-SGL-003: a single revision - trivially its own latest.
#: DOC-ASB-004: a single AS-BUILT ('Z') revision and nothing else - never
#: eligible to be latest, per `engine.revision.parsing`; kept here so the
#: exception path (no winner in the group) is exercised deliberately, not by
#: accident.
MIXED_DOCUMENT_ROWS: tuple[tuple[str, str], ...] = (
    ("DOC-NUM-001", "0"),
    ("DOC-NUM-001", "1"),
    ("DOC-NUM-001", "2"),
    ("DOC-ALP-002", "A"),
    ("DOC-ALP-002", "B"),
    ("DOC-SGL-003", "5"),
    ("DOC-ASB-004", "Z"),
)


def build_revision_workbook(path: Path,
                            document_rows: tuple[tuple[str, str], ...]
                            ) -> Path:
    """Write a QatarEnergy-TN-shaped workbook whose data rows are exactly
    `document_rows` (document number, revision), plus the companion sheets
    `MdrEngine.run()` reads."""
    wb = Workbook()

    ws = wb.active
    ws.title = "QatarEnergy-TN"
    for i, caption in enumerate(HEADERS, start=1):
        ws.cell(row=HEADER_ROW, column=i, value=caption)

    for r, (document_no, rev) in enumerate(document_rows, start=FIRST_DATA_ROW):
        ws.cell(row=r, column=1, value=document_no)
        ws.cell(row=r, column=2, value=rev)
        ws.cell(row=r, column=3, value="ELE")
        ws.cell(row=r, column=4, value=f"TITLE FOR {document_no}")
        ws.cell(row=r, column=5, value="IFR")
        ws.cell(row=r, column=6, value="")
        ws.cell(row=r, column=7, value="")
        ws.cell(row=r, column=8, value="SIGNED")

    vendors = wb.create_sheet("TN FROM VENDORS")
    vendors["A1"] = "VENDOR/ SUBCON"
    vendors["B1"] = "TRANSMITTAL NO."

    codes = wb.create_sheet("Status Codes")
    codes["A1"] = "ISSUE CODE"
    codes["B1"] = "REVIEW CODE FROM QatarEnergy"

    vendor_list = wb.create_sheet("VENDOR LIST")
    vendor_list["A1"] = "VENDOR"

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    wb.close()
    return path
